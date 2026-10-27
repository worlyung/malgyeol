"""녹취록 내보내기: TXT / MD / SRT / JSON / DOCX, 그리고 AI 요약 MD / DOCX."""
from __future__ import annotations

import io
import json
import re
import time
from typing import Iterable


def hms(t: float, ms: bool = False) -> str:
    t = max(0.0, float(t))
    h, rem = divmod(int(t), 3600)
    m, s = divmod(rem, 60)
    if ms:
        return f"{h:02d}:{m:02d}:{s:02d},{int((t % 1) * 1000):03d}"
    return f"{h:02d}:{m:02d}:{s:02d}"


def to_text(segments: Iterable[dict], title: str = "", timestamp: bool = True, speaker: bool = True) -> str:
    lines: list[str] = []
    if title:
        lines += [title, "=" * (len(title) + 4), ""]
    for s in segments:
        head = []
        if timestamp:
            head.append(f"[{hms(s['start'])}]")
        if speaker:
            head.append(f"{s.get('speaker') or '화자 미상'}:")
        lines.append(" ".join(head + [s["text"]]).strip())
    return "\n".join(lines) + "\n"


def info_lines(meta: dict | None) -> list[tuple[str, str]]:
    """머리말에 넣을 회의 정보. 적은 칸만 넣고, 일시를 비워 두면 녹음한 때를 씁니다."""
    if not meta:
        return []
    info = meta.get("info") or {}
    date = info.get("date") or (time.strftime("%Y-%m-%d %H:%M", time.localtime(meta["created"]))
                                if meta.get("created") else "")
    rows = [("일시", date), ("장소", info.get("place", "")), ("참석 인원", info.get("count", "")),
            ("참석자", info.get("attendees", ""))]
    return [(k, v) for k, v in rows if v]


def to_md(segments: Iterable[dict], title: str = "", timestamp: bool = True, speaker: bool = True,
          meta: dict | None = None) -> str:
    lines: list[str] = [f"# {title}", ""] if title else []
    head = info_lines(meta)
    if head:
        lines += [f"- **{k}**: {v}" for k, v in head] + ["", "---", ""]
    for s in segments:
        head = []
        if timestamp:
            head.append(f"[{hms(s['start'])}]")
        if speaker:
            head.append(f"{s.get('speaker') or '화자 미상'}:")
        lines += [(f"**{' '.join(head)}** " if head else "") + s["text"], ""]
    return "\n".join(lines)


def to_srt(segments: Iterable[dict], speaker: bool = True) -> str:
    out = []
    for i, s in enumerate(segments, 1):
        text = s["text"]
        if speaker and s.get("speaker"):
            text = f"{s['speaker']}: {text}"
        out.append(f"{i}\n{hms(s['start'], True)} --> {hms(s['end'], True)}\n{text}\n")
    return "\n".join(out)


def to_json(segments: list[dict], meta: dict) -> str:
    return json.dumps({"meta": meta, "segments": segments}, ensure_ascii=False, indent=2)


def _new_doc():  # noqa: ANN202
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10.5)
    # 한글은 eastAsia 글꼴을 따로 지정해야 워드에서 제대로 보입니다.
    from docx.oxml.ns import qn

    style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    return doc


def _docx_head(doc, title: str, meta: dict | None) -> None:  # noqa: ANN001
    if title:
        doc.add_heading(title, level=1)
    head = info_lines(meta)
    if head:
        table = doc.add_table(rows=len(head), cols=2)
        table.style = "Table Grid"
        for row, (k, v) in zip(table.rows, head):
            row.cells[0].text, row.cells[1].text = k, v
            row.cells[0].paragraphs[0].runs[0].bold = True
        doc.add_paragraph()


def to_docx(segments: list[dict], title: str = "", timestamp: bool = True, speaker: bool = True,
            meta: dict | None = None) -> bytes:
    from docx.shared import Pt

    doc = _new_doc()
    _docx_head(doc, title, meta)

    last_speaker = None
    for s in segments:
        p = doc.add_paragraph()
        name = s.get("speaker") or "화자 미상"
        if speaker and name != last_speaker:
            run = p.add_run(f"{name} ")
            run.bold = True
            last_speaker = name
        if timestamp:
            run = p.add_run(f"[{hms(s['start'])}] ")
            run.font.size = Pt(8)
        p.add_run(s["text"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def summary_md(summary: str, title: str = "") -> str:
    # 회의 정보는 AI가 요약의 '회의 개요'에 이미 넣으므로 제목만 붙입니다
    return (f"# {title} — 회의 요약\n\n" if title else "") + summary.strip() + "\n"


def summary_docx(summary: str, title: str = "") -> bytes:
    """AI 요약(마크다운: 제목·목록·굵게·구분선)을 워드 문서로 옮깁니다."""
    doc = _new_doc()
    if title:
        doc.add_heading(f"{title} — 회의 요약", level=1)

    def add_runs(p, text: str) -> None:  # noqa: ANN001
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            if part:
                p.add_run(part).bold = i % 2 == 1

    for line in summary.splitlines():
        line = line.rstrip()
        if not line.strip() or re.fullmatch(r"\s*(---+|\*\*\*+|===+)\s*", line):
            continue
        if m := re.match(r"^(#{1,4})\s+(.*)$", line):
            doc.add_heading(m.group(2).replace("**", ""), level=min(4, len(m.group(1)) + 1))
        elif m := re.match(r"^\s*[-*•]\s+(.*)$", line):
            add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := re.match(r"^\s*\d+[.)]\s+(.*)$", line):
            add_runs(doc.add_paragraph(style="List Number"), m.group(1))
        else:
            add_runs(doc.add_paragraph(), line.lstrip("> "))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def plain_transcript(segments: list[dict]) -> str:
    """요약 모델에 넣을 본문 (시각 + 화자 + 발언)."""
    return "\n".join(
        f"[{hms(s['start'])}] {s.get('speaker') or '화자 미상'}: {s['text']}" for s in segments
    )
